# -*- coding: utf-8 -*-
"""Génère les 3 documents Word Ayeba (rapport fondateur, guide équipe, guide partenaires)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"C:\Users\ADMIN\Documents\Ayeba"
os.makedirs(OUT, exist_ok=True)

ACCENT = RGBColor(0x0B, 0x5C, 0xD6)
GREY = RGBColor(0x55, 0x55, 0x55)


def base_doc(title, subtitle):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    t = doc.add_heading(title, level=0)
    for r in t.runs:
        r.font.color.rgb = ACCENT
    p = doc.add_paragraph(subtitle)
    p.runs[0].font.color.rgb = GREY
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Document généré le 21 mars 2026 — Ayeba / DevAlpha").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    return doc


def h1(doc, text):
    h = doc.add_heading(text, level=1)
    for r in h.runs:
        r.font.color.rgb = ACCENT


def h2(doc, text):
    h = doc.add_heading(text, level=2)


def para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


def bullets(doc, items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def numbered(doc, items):
    for i in items:
        doc.add_paragraph(i, style="List Number")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t


# =====================================================================
# DOC 1 — RAPPORT FONDATEUR (interne, complet, sans secrets)
# =====================================================================
doc = base_doc(
    "AYEBA — Rapport technique et produit",
    "Document interne — Fondateur · État complet de la plateforme",
)

h1(doc, "1. Résumé exécutif")
para(doc, "Ayeba est un écosystème web intégré destiné aux utilisateurs francophones et à la RDC, construit selon les standards des produits Google : moteur de recherche réel (crawler + index maison), console webmaster (Ayeba Studio), plateforme développeur OAuth 2.0/OpenID Connect (Ayeba Developers), encyclopédie collaborative (Ayebi), réseau publicitaire diffusé sur les applications sœurs (JEMSA, TALA, Omega, Sombateka), paiements réels (Stripe, CinetPay Mobile Money, Flutterwave), et back-office d'administration complet.")
para(doc, "Principe directeur appliqué sur tout le code : aucune donnée simulée. Chaque métrique affichée provient d'une table alimentée par du trafic, du crawl ou des transactions réels ; les ensembles vides affichent des états vides honnêtes plutôt que des chiffres inventés.", italic=True)

h1(doc, "2. État des produits")
table(doc,
    ["Produit", "Référence", "État", "Détail"],
    [
        ["Recherche Ayeba", "Google Search", "Production", "Crawler HTTP réel (robots.txt, extraction titres/liens/produits), index FTS5, signaux impressions/clics, suggestions, historique"],
        ["Ayeba Studio — Trace", "GA4 + Tag Manager", "Production", "Collecteur JS réel (sessions, devices, géo IP, UTM, gclid/fbclid), 6 pages : temps réel, audience, acquisition, comportement, conversions, balises"],
        ["Ayeba Studio — Radar", "Search Console", "Production", "5 pages : performance (requêtes/pages/pays/appareils), inspection d'URL, couverture, sitemaps, liens réels du crawler"],
        ["Ayeba Studio — Yield", "Google Ads", "Production", "Campagnes, créas avec revue, mots-clés, audiences à règles, facturation, placements éditeurs — données réelles"],
        ["Ayeba Studio — Velocity", "PageSpeed Insights", "Production", "Audits via API Google PSI réelle + fallback audit direct (fetch, TTFB, analyse HTML)"],
        ["Ayeba Studio — Aether", "Trends + Merchant Center", "Production", "Tendances issues de search_history réel, catalogue produits du crawler, diagnostics actionnables"],
        ["Ayeba Developers", "Cloud Console", "Production", "5 pages console : projets, clés API (restrictions + quotas), OAuth/OIDC avec scopes appliqués, métriques d'usage, journaux. API publique réelle /api/v1/search"],
        ["Ayebi", "Wikipédia", "Production", "Articles riches (sections, sous-sections, infobox, chronologie, galerie, références numérotées), édition par section, historique + diff, discussion, watchlist, protection, portails, backlinks"],
        ["Back-office /admin", "Console interne", "Production", "10 onglets : vue d'ensemble, utilisateurs, modération, support (tickets), facturation, réseau pubs, contenu, journal d'audit, équipe admin, chat interne"],
        ["Paiements", "Stripe + Mobile Money", "Production", "Stripe Checkout + webhooks signés, CinetPay (M-Pesa/Airtel/Orange) avec re-vérification serveur, Flutterwave (cartes, MoMo, payouts)"],
        ["Réseau pubs apps sœurs", "AdSense-like", "Production", "/api/ads/serve par clé d'app, tracking signé HMAC, anti-fraude (dédup, hash IP cryptographique)"],
    ])

h1(doc, "3. Architecture technique")
bullets(doc, [
    "Next.js App Router + TypeScript, déployé sur Vercel, domaine https://ayeba.app (~300 routes).",
    "Base SQLite en abstraction multi-mode : fichier local en dev, Turso (libsql) en production, mémoire pour les tests. Attention : les appels synchrones libsql sont interdits en prod (chemins async dédiés dans storage/turso-async).",
    "Schéma appliqué au démarrage + migrations de colonnes idempotentes (ALTER protégés).",
    "Crawler réel : requêtes HTTP, respect robots, extraction contenu/liens/images/produits, alimente index + Radar + Aether.",
    "OAuth 2.0 / OIDC complet : authorize + PKCE, codes, access/refresh tokens hashés, consentements, JWKS, scopes restreints par client, audit trail.",
])

h1(doc, "4. Sécurité")
bullets(doc, [
    "Sessions JWT signées, cookies httpOnly ; TOTP optionnel par compte.",
    "Clés API : format ayb_live_…, stockées en hash SHA-256, jamais rejouables en clair ; restrictions référent/IP, quotas journaliers, journalisation de chaque appel avec latence.",
    "Webhooks paiement : signatures vérifiées (Stripe HMAC + fenêtre anti-rejeu, Flutterwave verif-hash + re-confirmation API, CinetPay re-vérification serveur), idempotence via table payment_events.",
    "Pubs : clés par app sœur (comparaison timing-safe), CORS restreint, URLs de tracking signées HMAC (impressions/clics non falsifiables), déduplication.",
    "Admin : requireAdmin() central, permissions par rôle (super_admin/manager/support/moderator/analyst), dernier super_admin protégé, journal d'audit complet.",
    "Fail-closed : aucun provider sans clé ne simule un succès — erreur explicite.",
])

h1(doc, "5. Qualité et vérifications")
bullets(doc, [
    "Typecheck TypeScript : propre sur tout le graphe (~60 erreurs latentes corrigées dans les couches admin/studio/enterprise).",
    "ESLint : propre sur les nouveaux modules.",
    "Tests automatisés : 37 tests Vitest — webhooks Stripe (signature, rejeu), registry paiements, signatures HMAC pubs, permissions admin, clés API développeur (validation, quota, révocation, restrictions, isolation).",
    "Build Vercel : réussi, vérifié en production sur ayeba.app.",
])

h1(doc, "6. Problèmes majeurs corrigés")
bullets(doc, [
    "Couche Studio v2 jamais fonctionnelle : SQL dialecte PostgreSQL (EXTRACT) dans SQLite, colonnes inexistantes (position vs position_sum), table attribution_touchpoints manquante, paramètres bindés incorrects — toutes requêtes réparées et testées.",
    "Trace : chaque pageview créait une nouvelle session (ID client ignoré) et comptait double — corrigé ; métadonnées codées en dur (résolution/langue/fuseau) remplacées par les vraies valeurs client.",
    "Ad server : lecture de champs camelCase inexistants sur lignes snake_case (crash à l'exécution) — mappers ligne→domaine ajoutés.",
    "updateAdminUser écrivait une colonne updated_at inexistante — schéma corrigé.",
    "publishers.payment_methods (colonne inexistante) parsé → crash supprimé.",
])

h1(doc, "7. Ce qui reste à faire (honnête)")
bullets(doc, [
    "Les tableaux Studio afficheront des états vides tant qu'il n'y a pas de trafic réel — c'est voulu (jamais de faux chiffres).",
    "Velocity : configurer PAGESPEED_API_KEY pour activer les audits Lighthouse Google (sinon audit direct réel en fallback).",
    "CinetPay payouts : code prêt, attend les identifiants Transfer.",
    "Pas encore : email transactionnel (Resend), rate-limit distribué (Upstash), tests E2E navigateur.",
    "Variables à renseigner en production : clés Stripe/CinetPay/Flutterwave, AD_KEY_<APP> par application sœur, clés OAuth sœurs.",
])

doc.save(os.path.join(OUT, "AYEBA_Rapport_Fondateur.docx"))
print("1/3 rapport fondateur OK")


# =====================================================================
# DOC 2 — GUIDE ÉQUIPE (collaborateurs / travailleurs)
# =====================================================================
doc = base_doc(
    "AYEBA — Guide de l'équipe",
    "Mode d'emploi interne — modération, support, contribution, outils",
)

h1(doc, "1. Accès")
numbered(doc, [
    "Créez votre compte Ayeba sur https://ayeba.app/ayebi/connexion (mode inscription).",
    "Le compte seul ne donne AUCUN accès : le super administrateur vous ajoute dans l'équipe (onglet « Équipe admin ») et vous transmet un mot de passe provisoire si nécessaire.",
    "Ouvrez https://ayeba.app/admin — vous ne voyez que les onglets correspondant à votre rôle.",
])

table(doc,
    ["Rôle", "Ce que vous voyez", "Ce que vous pouvez faire"],
    [
        ["super_admin", "Tous les onglets", "Tout, y compris gérer l'équipe"],
        ["manager", "Vue d'ensemble, modération, support, facturation, réseau pubs, contenu, chat", "Valider campagnes, gérer éditeurs, lire billing, modérer, assigner tickets"],
        ["support", "Vue d'ensemble, modération (lecture), support, chat", "Répondre aux tickets, changer leur statut"],
        ["moderator", "Vue d'ensemble, modération, contenu, chat", "Approuver/rejeter dans la file de modération"],
        ["analyst", "Vue d'ensemble, chat", "Lecture seule"],
    ])

h1(doc, "2. Support utilisateurs (onglet Support)")
numbered(doc, [
    "Ouvrez un ticket dans la liste — la conversation complète s'affiche.",
    "Rédigez la réponse, changez le statut (ouvert → en cours → résolu → fermé).",
    "Un manager peut assigner un ticket à un membre de l'équipe.",
    "Règle d'or : ton professionnel, réponse en français, jamais de données personnelles demandées par ticket.",
])

h1(doc, "3. Modération (onglet Modération)")
numbered(doc, [
    "La file affiche les contenus signalés (articles Ayebi, annonces, etc.).",
    "Approuver = le contenu reste publié ; Rejeter = retrait avec motif obligatoire.",
    "Chaque décision est écrite dans le journal d'audit avec votre nom — soyez rigoureux.",
])

h1(doc, "4. Chat d'équipe (onglet Chat équipe)")
bullets(doc, [
    "Canal « # équipe » visible par tous les administrateurs.",
    "Messages directs possibles admin ↔ admin via la liste des membres.",
    "Utilisez-le pour la coordination quotidienne — pas pour les informations confidentielles de paiement.",
])

h1(doc, "5. Contribuer à Ayebi (encyclopédie)")
numbered(doc, [
    "Créer : /ayebi/nouveau — titre, catégorie, résumé introductif, sections séparées par ---, sous-sections par ===.",
    "Modifier un article complet : bouton « Modifier » sur l'article.",
    "Modifier UNE section : lien [modifier] à côté du titre de la section — seul ce bloc est chargé dans l'éditeur.",
    "Infobox : lignes « Label|Valeur ». Chronologie : lignes « Date|Événement ». Références : syntaxe [ref:url|titre] dans le texte.",
    "Liens internes : [[slug|Texte affiché]]. L'aperçu montre le rendu avant publication.",
    "TOUJOURS remplir le résumé de modification — il apparaît dans l'historique public.",
    "Onglets utiles : Historique (diff entre versions, restauration), Discussion, Suivre (watchlist).",
])

h1(doc, "6. Principes non négociables")
bullets(doc, [
    "Jamais de fausses données : n'inventez pas de métriques, d'articles pompeux sans source, ni de statistiques.",
    "Toutes vos actions d'administration sont tracées dans le journal d'audit.",
    "Un doute sur une modération ou un remboursement → escaladez au super admin via le chat, ne décidez pas seul.",
    "Les mots de passe provisoires se transmettent en privé, jamais dans le canal public.",
])

doc.save(os.path.join(OUT, "AYEBA_Guide_Equipe.docx"))
print("2/3 guide equipe OK")


# =====================================================================
# DOC 3 — GUIDE PARTENAIRES / INVESTISSEURS (mode d'utilisation)
# =====================================================================
doc = base_doc(
    "AYEBA — Présentation & mode d'utilisation",
    "Document partenaires et investisseurs — l'écosystème numérique francophone",
)

h1(doc, "1. Qu'est-ce qu'Ayeba ?")
para(doc, "Ayeba est le premier écosystème numérique intégré conçu pour la RDC et l'espace francophone, sur le modèle de ce que Google est pour le web mondial : un moteur de recherche, une plateforme pour les webmasters, une plateforme pour les développeurs, une encyclopédie collaborative, et un réseau publicitaire diffusé sur une famille d'applications (JEMSA, TALA, Omega, Sombateka).")
para(doc, "Différence fondamentale avec une démo ou un prototype : tout est réel. Le moteur crawle véritablement le web, les statistiques proviennent de données collectées, les paiements passent par de vrais prestataires (Stripe, CinetPay Mobile Money, Flutterwave), et l'API développeurs répond avec de vrais résultats.", bold=True)

h1(doc, "2. Les produits, et comment les utiliser")

h2(doc, "2.1 Ayeba Search — le moteur de recherche")
bullets(doc, [
    "Accès : https://ayeba.app — tapez une requête, obtenez des résultats indexés par notre propre crawler.",
    "Côté technique : index maison (FTS5), signaux de qualité réels (impressions, clics, positions), suggestions instantanées.",
])

h2(doc, "2.2 Ayebi — l'encyclopédie libre (type Wikipédia)")
bullets(doc, [
    "Accès : https://ayeba.app/ayebi — articles structurés : infobox, sections, références numérotées, chronologie, galerie, cartes.",
    "Collaboratif : tout compte peut créer/modifier, avec historique complet, comparaison de versions, pages de discussion et modération.",
    "Focus : contenus sur la RDC et l'Afrique francophone — un savoir local que Wikipédia couvre peu.",
])

h2(doc, "2.3 Ayeba Studio — la console des webmasters (type Google Search Console + Analytics + Ads)")
bullets(doc, [
    "Accès : https://ayeba.app/studio — ajoutez votre site, prouvez la propriété, obtenez une clé de tracking.",
    "Trace : analytics complet (temps réel, audience, acquisition, comportement, conversions) + gestionnaire de balises.",
    "Radar : performance dans la recherche Ayeba (requêtes, pages, pays, appareils), inspection d'URL, couverture d'index, sitemaps, liens.",
    "Yield : publicité — campagnes, annonces, mots-clés, audiences, facturation. Vos annonces sont diffusées sur tout l'écosystème.",
    "Velocity : audits de performance réels (PageSpeed Insights).",
    "Aether : tendances de recherche réelles + catalogue produits détecté par le crawler.",
])

h2(doc, "2.4 Ayeba Developers — la plateforme développeurs (type Google Cloud Console)")
numbered(doc, [
    "Créez un compte, ouvrez https://ayeba.app/developers/console.",
    "Créez un projet, puis une clé API — la clé n'est affichée qu'une fois (format ayb_live_…).",
    "Appelez l'API : GET https://ayeba.app/api/v1/search?q=votre+requete avec l'en-tête Authorization: Bearer VOTRE_CLE.",
    "Définissez quotas journaliers et restrictions (domaines référents, IPs) par clé.",
    "OAuth 2.0/OpenID Connect : créez un client pour « Se connecter avec Ayeba » — scopes configurables, consentement utilisateur, PKCE.",
    "Suivez votre usage : appels par jour, erreurs, latence, journaux détaillés.",
])

h2(doc, "2.5 Publicité sur l'écosystème")
bullets(doc, [
    "Les annonces créées dans Yield sont servies sur Ayeba et les applications sœurs JEMSA, TALA, Omega et Sombateka.",
    "Tracking signé et anti-fraude : chaque impression/clic est authentifié et dédupliqué — vous payez pour du trafic réel.",
    "Facturation : factures et transactions réelles consultables dans la console.",
])

h2(doc, "2.6 Paiements")
bullets(doc, [
    "Mobile Money RDC via CinetPay (M-Pesa, Airtel Money, Orange Money).",
    "Cartes bancaires via Stripe et Flutterwave ; payouts éditeurs via Flutterwave Transfers.",
    "Sécurité bancaire : webhooks signés, re-vérification serveur des statuts, anti-double-crédit.",
])

h1(doc, "3. Modèle économique")
bullets(doc, [
    "Publicité : annonceurs facturés via Yield (enchères CPC/CPM), revenus partagés avec les éditeurs et apps sœurs.",
    "API développeurs : quotas gratuits puis paliers payants (infrastructure déjà en place).",
    "Services Studio premium : audits, rapports, visibilité accrue.",
    "Position unique : aucun acteur ne couvre aujourd'hui la RDC/francophonie avec un stack intégré recherche + contenu + publicité + identité.",
])

h1(doc, "4. Chiffres clés techniques")
table(doc,
    ["Indicateur", "Valeur"],
    [
        ["Routes applicatives déployées", "~300"],
        ["Domaine de production", "https://ayeba.app"],
        ["Couverture tests automatisés", "37 tests (paiements, sécurité, pubs, admin, API)"],
        ["Providers paiement", "Stripe, CinetPay, Flutterwave"],
        ["Applications sœurs connectables", "JEMSA, TALA, Omega, Sombateka"],
        ["Données simulées", "0 — toutes les métriques proviennent de données réelles"],
    ])

h1(doc, "5. Feuille de route")
numbered(doc, [
    "Densification de l'index de recherche (crawl continu du web francophone et congolais).",
    "Croissance du corpus Ayebi avec contributeurs locaux.",
    "Ouverture commerciale de Yield aux annonceurs et éditeurs.",
    "Applications sœurs branchées sur l'identité Ayeba (OAuth) et le réseau publicitaire.",
    "Notifications email transactionnelles et extension de l'API publique (nouveaux endpoints).",
])

para(doc, "")
para(doc, "Contact : benjaminoussama62@gmail.com — Ayeba / DevAlpha", italic=True)

doc.save(os.path.join(OUT, "AYEBA_Guide_Partenaires_Investisseurs.docx"))
print("3/3 guide partenaires OK")

print("\nFichiers dans " + OUT)
